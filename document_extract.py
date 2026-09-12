import csv
import json
import os
import sys
from pathlib import Path


def configure_tesseract(pytesseract):
    command = os.environ.get('TESSERACT_CMD')
    windows_command = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
    if command and os.path.exists(command):
        pytesseract.pytesseract.tesseract_cmd = command
    elif os.path.exists(windows_command):
        pytesseract.pytesseract.tesseract_cmd = windows_command


def extract(path):
    suffix = Path(path).suffix.lower()
    if suffix in {'.txt', '.md', '.log', '.xml', '.html'}:
        return Path(path).read_text(encoding='utf-8', errors='replace')
    if suffix == '.json':
        return json.dumps(json.loads(Path(path).read_text(encoding='utf-8')), indent=2, ensure_ascii=False)
    if suffix == '.csv':
        with open(path, newline='', encoding='utf-8-sig', errors='replace') as handle:
            return '\n'.join(' | '.join(row) for row in csv.reader(handle))
    if suffix == '.pdf':
        from pypdf import PdfReader
        text = '\n\n'.join(page.extract_text() or '' for page in PdfReader(path).pages).strip()
        if text:
            return text
        import pymupdf
        from PIL import Image
        import pytesseract
        configure_tesseract(pytesseract)
        document = pymupdf.open(path)
        pages = []
        for page in document:
            pixels = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            image = Image.frombytes('RGB', [pixels.width, pixels.height], pixels.samples)
            pages.append(pytesseract.image_to_string(image).strip())
        return '\n\n'.join(pages).strip()
    if suffix == '.docx':
        from docx import Document
        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            parts.append('\n'.join(' | '.join(cell.text.strip() for cell in row.cells) for row in table.rows))
        for section in document.sections:
            parts.extend(paragraph.text for paragraph in section.header.paragraphs if paragraph.text.strip())
            parts.extend(paragraph.text for paragraph in section.footer.paragraphs if paragraph.text.strip())
        return '\n'.join(parts)
    if suffix in {'.xlsx', '.xlsm'}:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for sheet in workbook.worksheets:
            rows = [' | '.join('' if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True)]
            sheets.append(f'[{sheet.title}]\n' + '\n'.join(rows))
        return '\n\n'.join(sheets)
    raise ValueError(f'Unsupported file type: {suffix or "unknown"}')


try:
    print(json.dumps({'text': extract(sys.argv[1])}, ensure_ascii=False))
except Exception as error:
    print(json.dumps({'error': str(error)}))
    sys.exit(1)