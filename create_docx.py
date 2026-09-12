import json
import re
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def clean_inline(value):
    return re.sub(r'[*_`]', '', value).strip()


def table_cells(line):
    return [clean_inline(cell) for cell in line.strip().strip('|').split('|')]


def is_table_separator(line):
    return bool(re.match(r'^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$', line))


def add_table(document, rows):
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = 'Table Grid'
    for cell, value in zip(table.rows[0].cells, rows[0]):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for values in rows[1:]:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value


def add_page_numbers(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    run._r.addnext(field)


def add_content(document, content):
    lines = content.splitlines()
    index = 0
    while index < len(lines):
        raw_line = lines[index]
        line = raw_line.strip()
        if not line:
            index += 1
            continue
        if line.startswith('|') and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            rows = [table_cells(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append(table_cells(lines[index]))
                index += 1
            width = len(rows[0])
            add_table(document, [row + [''] * (width - len(row)) for row in rows])
            continue
        if line.startswith('### '):
            document.add_heading(clean_inline(line[4:]), level=3)
        elif line.startswith('## '):
            document.add_heading(clean_inline(line[3:]), level=2)
        elif line.startswith('# '):
            document.add_heading(clean_inline(line[2:]), level=1)
        elif re.match(r'^[-*] ', line):
            document.add_paragraph(clean_inline(line[2:]), style='List Bullet')
        elif re.match(r'^\d+\. ', line):
            document.add_paragraph(clean_inline(re.sub(r'^\d+\. ', '', line)), style='List Number')
        elif line == '---':
            document.add_paragraph()
        else:
            document.add_paragraph(clean_inline(line))
        index += 1


with open(sys.argv[1], encoding='utf-8-sig') as source:
    data = json.load(source)

document = Document()
section = document.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
normal = document.styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
normal.font.size = Pt(12)
for style_name, size in [('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 12)]:
    style = document.styles[style_name]
    style.font.name = 'Times New Roman'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    style.font.size = Pt(size)
    style.font.bold = True
add_page_numbers(section)
title = document.add_heading(data.get('title') or 'Engibot Document', level=0)
title.runs[0].font.name = 'Times New Roman'
title.runs[0].font.size = Pt(18)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_content(document, data.get('content', ''))
document.save(data['output'])